"""Verify the derived second contract, empty actuals and the real two-contract aggregation."""
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from compare_excel_app import load_version, normalized

model = ROOT / 'data/contractors/c1/models/model-2026-09-22.xlsx'
original, pilot = load_version(model, 'd_pilot_4700134128', 'v_pilot_2026')
scenario, second = load_version(model, 'd_scenario_b_4700134128', 'v_scenario_b_2026')
_, second_next = load_version(model, 'd_scenario_b_4700134128', 'v_scenario_b_2027')
folder = ROOT / 'data/contractors/c1/sources/scenario-b'
boq = openpyxl.load_workbook(folder / 'boq.xlsx', read_only=True, data_only=True).active
ksg = openpyxl.load_workbook(folder / 'ksg.xlsx', read_only=True, data_only=True).active
original_ksg = openpyxl.load_workbook(ROOT / 'data/sources/ksg.xlsx', read_only=True, data_only=True).active
ksg_rows = list(ksg.values)
original_ksg_rows = list(original_ksg.values)
codes = {row[31] for row in boq.values if len(row) > 31 and isinstance(row[31], str) and row[31].startswith('KQ-')}
assert 50 <= len(codes) <= 75
assert scenario['number'] in str(boq['U2'].value)
assert second['sourceFiles']['boq'].endswith('scenario-b/boq.xlsx')
assert second['sourceFiles']['ksg'].endswith('scenario-b/ksg.xlsx')
assert second['sourceFiles']['primary_documents'].endswith('data/sources/primary_documents.xlsx')
assert scenario['sourceIdentities']['primary_documents']['contractNumber'] == scenario['number']
assert scenario['sourceIdentities']['payments']['contractNumber'] == scenario['number']
primary = openpyxl.load_workbook(ROOT / 'data/sources/primary_documents.xlsx', read_only=True, data_only=True).active
payments = openpyxl.load_workbook(ROOT / 'data/sources/payments.xlsx', read_only=True, data_only=True).active
assert not any(row[1] == 'Подрядчик-1--2' and row[5] == scenario['number'] and row[6] == 'ПЭ'
               for row in primary.values if len(row) > 6)
assert not any(str(row[4]) == scenario['number'] and scenario['number'] in str(row[12])
               for row in payments.values if len(row) > 12)
assert second['actualThroughMonth'] == second_next['actualThroughMonth'] == 0
for version in (second, second_next):
    assert not any(version['primaryMonthsWithDocuments'])
    assert not any(version['paymentActualMonths'])
    assert not any(version['drivers']['primaryExecuted'])
    assert not any(version['drivers']['payments'])
    assert version['workSourceMeta']['unmatchedRows'] == 0
    assert version['workItems'] and sum(version['drivers']['ksgRevenue']) > 0
assert sum(second_next['drivers']['ksgRevenue']) > 0

shifted = 0
for row in range(5, 171):
    code = ksg_rows[row - 1][1] if len(ksg_rows[row - 1]) > 1 else None
    if code not in codes:
        continue
    start = ksg_rows[row - 1][14]
    original_start = original_ksg_rows[row - 1][14]
    assert str(start)[:10] == str(original_start)[:10]
    for col in range(19, 42):
        old = original_ksg_rows[row - 1][col - 1] or 0
        new = ksg_rows[row - 1][col + 2] or 0
        assert abs(float(old) - float(new)) < 1e-8, (row, col)
    shifted += 1
assert shifted > 50

ratio = second['workSourceMeta']['scenarioProvenance']['personnelScale']
assert .25 < ratio < .75
assert abs(second['drivers']['directPeople'][3] - pilot['drivers']['directPeople'][0] * ratio) < 1e-8
assert abs(second_next['drivers']['directPeople'][0] - pilot['drivers']['directPeople'][9] * ratio) < 1e-8

entries = [{'contractId': contract['id'], 'number': contract['number'],
            'version': normalized(version, contract, ROOT)}
           for contract, version in ((original, pilot), (scenario, second))]
run = subprocess.run(['node', str(ROOT / 'scripts/portfolio_snapshot.js')],
                     input=json.dumps(entries), text=True, capture_output=True, check=True)
snapshot = json.loads(run.stdout)
assert snapshot['status'] == 'ready' and snapshot['year'] == 2026
assert snapshot['currency'] == 'USD' and len(snapshot['lines']) == 2
assert all(abs(x) < 1e-6 for x in snapshot['control'])
for key in ('revenue', 'costs', 'ncf'):
    assert all(abs(snapshot['totals'][key][m] - sum(row[key][m] for row in snapshot['individual'])) < 1e-5
               for m in range(12)), key
assert any('расчётным сценарием' in warning for warning in snapshot['warnings'])
pilot_result, scenario_result = snapshot['individual']
pilot_fact = max(i + 1 for i, x in enumerate(pilot['paymentActualMonths']) if x)
assert pilot_result['receipts'][:pilot_fact] == pilot['drivers']['payments'][:pilot_fact], 'bank fact preserved'
assert any(value > 0 for value in pilot_result['receipts'][pilot_fact:]), 'pilot forecast payments'
assert any(value > 0 for value in scenario_result['receipts']), 'scenario cash forecast from KSG/KS-2'
assert all(abs(x) < 1e-5 for x in scenario_result['receivable'][-3:]), 'forecast closes credit balance'
print('TWO CONTRACTS: OK; KSG rows', shifted, 'personnel scale', round(ratio, 4))
