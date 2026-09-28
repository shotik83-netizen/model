#!/usr/bin/env python3
"""Add the archived visual examples to a copy of the current pilot workbook."""

import io
import json
import subprocess
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'data/contractors/c1/models/model-2026-09-22.xlsx'
OUTPUT = ROOT / 'data/contractors/c1/models/model-2026-09-28-demo.xlsx'
ARCHIVE = '390822c^:data/contractors/c1/models/model-2026-09-22.xlsx'


def payload(book):
    sheet = book['_DATA']
    return json.loads(''.join(str(row[0] or '') for row in list(sheet.values)[1:]))


base_book = load_workbook(BASE)
archived_bytes = subprocess.check_output(['git', 'show', ARCHIVE], cwd=ROOT)
archived = payload(load_workbook(io.BytesIO(archived_bytes), read_only=True))
current = payload(base_book)
contracts = current['contractor']['contracts']
pilot = next(c for c in contracts if c['id'] == 'd_pilot_4700134128')
archived_pilot = next(c for c in archived['contractor']['contracts'] if c['id'] == pilot['id'])
factor = next(v for v in archived_pilot['versions'] if v['id'] == 'v_factor_example_2026')
assert [v['id'] for v in pilot['versions']] == ['v_pilot_2026']
assert factor['exampleEdits']
pilot['versions'].insert(0, factor)
scenario = next(c for c in archived['contractor']['contracts'] if c['id'] == 'd_scenario_b_4700134128')
assert all(v.get('scenarioDerived') for v in scenario['versions'])
contracts.append(scenario)

raw = json.dumps(current, ensure_ascii=False, separators=(',', ':'))
sheet = base_book['_DATA']
sheet.delete_rows(2, sheet.max_row)
for index in range(0, len(raw), 30000):
    sheet.cell(2 + index // 30000, 1, raw[index:index + 30000])
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
base_book.save(OUTPUT)
print(f'RESTORED DEMO MODELS: {OUTPUT.relative_to(ROOT)}')
