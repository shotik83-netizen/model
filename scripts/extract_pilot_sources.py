"""Read-only extraction of source workbook cells for the app's BOQ/KSG importer."""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


def cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def rows(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    result = []
    for line in sheet.iter_rows(values_only=True):
        values = [cell(v) for v in line]
        while values and values[-1] is None:
            values.pop()
        result.append(values)
    return {'name': sheet.title, 'rows': result, 'maxRow': len(result)-1}


def resource_months(path, target_rows):
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    selected = {}
    for index, line in enumerate(sheet.iter_rows(min_row=4,max_row=max(target_rows),values_only=True),4):
        if index in {4, *target_rows}:
            selected[index] = line
    dates = selected[4]
    result = {}
    for row in target_rows:
        series = []
        for month in range(1,13):
            values = [v for d,v in zip(dates,selected[row])
                      if isinstance(d,(date,datetime)) and d.year==2026 and d.month==month]
            series.append({'days':len(values),'filledDays':sum(v is not None for v in values),
                           'average':sum(float(v or 0) for v in values)/len(values) if values else 0})
        result[str(row)] = series
    return result


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1] / 'data/sources'
    output = Path(sys.argv[1])
    data={key: {'sheets': [rows(root / (key + '.xlsx'))]}
          for key in ['boq', 'ksg', 'payments', 'primary_documents', 'factoring']}
    data['resourceMonths']={'personnel':resource_months(root/'personnel.xlsx',[325,326,327]),
                            'equipment':resource_months(root/'equipment.xlsx',[227])}
    output.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')))
    print(output, output.stat().st_size)
